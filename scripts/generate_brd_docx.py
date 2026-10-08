"""Generate an editable Word BRD from the canonical Markdown document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "CLOUDSENTINEL_CSPM_BRD.md"
OUTPUT = ROOT / "docs" / "CLOUDSENTINEL_CSPM_BRD.docx"
NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = clean_inline(row[column_index]) if column_index < len(row) else ""
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document: Document, *, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    for level, size, color in ((1, 18, NAVY), (2, 14, NAVY), (3, 11, BLUE)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if level > 1 else 14)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "CloudSentinel CSPM | Functional Business Requirements Document"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(BLUE)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core = document.core_properties
    core.title = title
    core.subject = "CSPM business and product requirements"
    core.author = "CloudSentinel Product Team"
    core.keywords = "CSPM, cloud security, BRD, AWS, Azure, GCP, IaC"


def build_document(
    source: Path = SOURCE,
    output: Path = OUTPUT,
    subtitle: str = "Market-informed product requirements and delivery gates",
) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = next(
        (clean_inline(line.removeprefix("# ")) for line in lines if line.startswith("# ")),
        "CloudSentinel CSPM Business Requirements Document",
    )
    document = Document()
    configure_document(document, title=title)
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    paragraph_buffer: list[str] = []
    figure_number = 0

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        paragraph = document.add_paragraph(clean_inline(" ".join(paragraph_buffer)))
        paragraph.paragraph_format.keep_together = True
        paragraph_buffer.clear()

    def flush_table() -> None:
        if not table_rows:
            return
        filtered = [
            row
            for row in table_rows
            if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row)
        ]
        add_table(document, filtered)
        table_rows.clear()

    for line in lines:
        stripped = line.strip()
        code_fence = re.match(r"^```([A-Za-z0-9_-]*)", stripped)
        if code_fence:
            flush_paragraph()
            flush_table()
            if in_code:
                # Mermaid source remains in the editable Markdown. The Word
                # version embeds the corresponding rendered figure instead of
                # a non-editable text diagram.
                if code_language != "mermaid":
                    paragraph = document.add_paragraph("\n".join(code_lines))
                    paragraph.style = document.styles["No Spacing"]
                    set_cell = paragraph._p.get_or_add_pPr()
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), LIGHT_GRAY)
                    set_cell.append(shading)
                    for run in paragraph.runs:
                        run.font.name = "Consolas"
                        run.font.size = Pt(8.5)
                code_lines.clear()
                code_language = ""
                in_code = False
            else:
                in_code = True
                code_language = code_fence.group(1).lower()
            continue
        if in_code:
            code_lines.append(line)
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_rows.append([cell.strip() for cell in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped:
            flush_paragraph()
            continue
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image_match:
            flush_paragraph()
            image_path = Path(image_match.group(2))
            if not image_path.is_absolute():
                image_path = source.parent / image_path
            if image_path.exists():
                figure_number += 1
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_together = True
                paragraph.add_run().add_picture(str(image_path), width=Inches(7.0))
                caption = document.add_paragraph(
                    f"Figure {figure_number}. {clean_inline(image_match.group(1))}"
                )
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_after = Pt(6)
                for run in caption.runs:
                    run.italic = True
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string(BLUE)
            else:
                paragraph = document.add_paragraph(
                    f"[Image unavailable: {clean_inline(image_match.group(1))}]"
                )
                for run in paragraph.runs:
                    run.italic = True
                    run.font.color.rgb = RGBColor(192, 0, 0)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = clean_inline(heading.group(2))
            paragraph = document.add_heading(title, level=level)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                subtitle_paragraph = document.add_paragraph(subtitle)
                subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                subtitle_paragraph.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
                subtitle_paragraph.runs[0].italic = True
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            paragraph = document.add_paragraph(clean_inline(stripped.lstrip("> ")))
            paragraph.style = document.styles["Quote"]
            paragraph.paragraph_format.left_indent = Inches(0.3)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string(NAVY)
                run.bold = True
            continue
        unordered = re.match(r"^-\s+(.+)$", stripped)
        if unordered:
            flush_paragraph()
            document.add_paragraph(clean_inline(unordered.group(1)), style="List Bullet")
            continue
        ordered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if ordered:
            flush_paragraph()
            document.add_paragraph(clean_inline(ordered.group(1)), style="List Number")
            continue
        if stripped == "---":
            flush_paragraph()
            continue
        paragraph_buffer.append(stripped)

    flush_paragraph()
    flush_table()

    approval_heading = next(
        (p for p in document.paragraphs if p.text.strip().endswith(". Approval")),
        None,
    )
    if approval_heading:
        approval_heading.paragraph_format.page_break_before = True

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Generated {output} ({output.stat().st_size:,} bytes)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument(
        "--subtitle",
        default="Market-informed product requirements and delivery gates",
    )
    args = parser.parse_args()
    build_document(args.source, args.output, args.subtitle)
